from io import BytesIO
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from datetime import datetime
import os

# Default slots (used if no time_slots provided)
DEFAULT_TIME_SLOTS = [
    "10:00 am to 11:00 am",
    "11:00 am to 12:00 pm",
    "12:00 pm to 1:00 pm",
    "1:00 pm to 2:00 pm",
    "2:00 pm to 3:00 pm",
    "3:00 pm to 4:00 pm",
    "4:00 pm to 5:00 pm",
    "5:00 pm to 6:00 pm"
]

# Helper: Format date to DD-MM-YYYY
def format_date_for_excel(date_val):
    if not date_val:
        return ""
    if isinstance(date_val, str):
        try:
            dt = datetime.strptime(date_val, "%Y-%m-%d")
            return dt.strftime("%d-%m-%Y")
        except ValueError:
            return date_val
    elif isinstance(date_val, datetime):
        return date_val.strftime("%d-%m-%Y")
    return str(date_val)

def create_excel_bytes(schedule_data, template_bytes=None, time_slots=None):
    if time_slots is None:
        time_slots = DEFAULT_TIME_SLOTS

    wb = None
    ws = None

    # 1. Try to load template
    if template_bytes is None:
        template_path = os.path.join(os.path.dirname(__file__), '..', '..', 'EOD_SAMPLE_FINAL.xlsx')
        if os.path.exists(template_path):
            with open(template_path, 'rb') as f:
                template_bytes = f.read()
            print(f"✅ Loaded template from {template_path}")

    if template_bytes:
        try:
            wb = load_workbook(BytesIO(template_bytes))
            ws = wb.active
            print("✅ Template loaded successfully")
        except Exception as e:
            print(f"⚠️ Failed to load template: {e} – using built-in")
            wb = None

    # 2. Built-in layout if no template
    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = "EOD Report"

        ws.merge_cells('B1:C1')
        cell = ws['B1']
        cell.value = "EOD REPORT"
        cell.font = Font(name='Calibri', size=16, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')

        ws['B3'] = "Name Of Employee"
        ws['C3'] = ""
        ws['B4'] = "Position"
        ws['C4'] = ""
        ws['B5'] = "Date"
        ws['C5'] = ""

        for r in [3,4,5]:
            ws[f'B{r}'].font = Font(bold=True)

        ws['B7'] = "Time"
        ws['C7'] = "Activity"
        ws['D7'] = "Description"
        for col in ['B','C','D']:
            ws[f'{col}7'].font = Font(bold=True)

        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 28
        ws.column_dimensions['D'].width = 35

        thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                      top=Side(style='thin'), bottom=Side(style='thin'))
        for row in ws.iter_rows(min_row=7, max_row=len(time_slots)+7, min_col=2, max_col=4):
            for cell in row:
                cell.border = thin
        for cell_addr in ['C3','C4','C5']:
            ws[cell_addr].border = thin

    # 3. Fill data with formatted date
    try:
        ws['C3'] = schedule_data.get("employee_name", "")
        ws['C4'] = schedule_data.get("position", "")
        ws['C5'] = format_date_for_excel(schedule_data.get("date", ""))
    except Exception as e:
        print(f"⚠️ Could not fill employee details: {e}")

    # 4. Fill schedule rows
    start_row = 8
    for idx, slot in enumerate(time_slots):
        row = start_row + idx
        entry = next((item for item in schedule_data.get("schedule", []) if item.get("slot") == slot), None)
        activity = entry.get("activity", "") if entry else ""
        description = entry.get("description", "") if entry else ""

        if "Lunch Break" in slot or "lunch" in slot.lower():
            activity = "Lunch Break"
            description = "Lunch Break"

        try:
            ws[f'C{row}'] = activity
            ws[f'D{row}'] = description
        except Exception as e:
            print(f"⚠️ Could not fill row {row}: {e}")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
import subprocess
import tempfile
import os
from io import BytesIO
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from openpyxl import load_workbook
from .excel_service import create_excel_bytes

# -------- Helper: Format date to DD-MM-YYYY --------
def format_date_pdf(date_val):
    """
    Convert a date value to DD-MM-YYYY format.
    Accepts string (YYYY-MM-DD), datetime object, or None.
    """
    if not date_val:
        return ""
    if isinstance(date_val, str):
        # Try to parse YYYY-MM-DD
        try:
            dt = datetime.strptime(date_val, "%Y-%m-%d")
            return dt.strftime("%d-%m-%Y")
        except ValueError:
            # Maybe it's already formatted? Return as is.
            return date_val
    elif isinstance(date_val, datetime):
        return date_val.strftime("%d-%m-%Y")
    return str(date_val)

# -------- Helper: Find LibreOffice --------
def find_soffice():
    possible_paths = [
        "soffice",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for cmd in possible_paths:
        try:
            subprocess.run([cmd, "--version"], capture_output=True, check=True, timeout=5)
            return cmd
        except:
            continue
    return None

# -------- Main entry --------
def create_pdf_bytes(schedule_data, excel_bytes=None):
    # 1. If excel_bytes not provided, generate them from schedule_data
    if excel_bytes is None:
        excel_bytes = create_excel_bytes(schedule_data)

    # 2. Try LibreOffice conversion (produces exact replica)
    try:
        return _create_pdf_with_libreoffice(excel_bytes)
    except Exception as e:
        print(f"⚠️ LibreOffice failed: {e} – using template-driven fallback")
        return _create_pdf_with_template_fallback(schedule_data)

# -------- LibreOffice conversion (takes Excel bytes) --------
def _create_pdf_with_libreoffice(excel_bytes):
    soffice_cmd = find_soffice()
    if soffice_cmd is None:
        raise RuntimeError("LibreOffice not found")

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_xlsx:
        tmp_xlsx.write(excel_bytes)
        xlsx_path = tmp_xlsx.name

    pdf_path = xlsx_path.replace(".xlsx", ".pdf")

    try:
        subprocess.run([
            soffice_cmd, "--headless", "--convert-to", "pdf",
            "--outdir", os.path.dirname(pdf_path), xlsx_path
        ], check=True, capture_output=True, timeout=60)
        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()
        print("✅ PDF created with LibreOffice")
    finally:
        if os.path.exists(xlsx_path):
            os.unlink(xlsx_path)
        if os.path.exists(pdf_path):
            os.unlink(pdf_path)

    return pdf_bytes

# -------- Template-driven fallback (reads your Excel template) --------
def _create_pdf_with_template_fallback(schedule_data):
    # Locate the template
    template_path = os.path.join(os.path.dirname(__file__), '..', '..', 'EOD_SAMPLE_FINAL.xlsx')
    if not os.path.exists(template_path):
        template_path = "/app/EOD_SAMPLE_FINAL.xlsx"
        if not os.path.exists(template_path):
            print("⚠️ Template not found – using basic fallback")
            return _create_basic_fallback_pdf(schedule_data)

    try:
        wb = load_workbook(template_path, data_only=True)
        ws = wb.active

        # Extract layout
        title = ws['B1'].value or "EOD REPORT"
        emp_label = ws['B3'].value or "Name Of Employee"
        pos_label = ws['B4'].value or "Position"
        date_label = ws['B5'].value or "Date"
        time_header = ws['B7'].value or "Time"
        act_header = ws['C7'].value or "Activity"
        desc_header = ws['D7'].value or "Description"

        # Prepare PDF
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4,
                                rightMargin=15, leftMargin=15,
                                topMargin=15, bottomMargin=15)
        elements = []
        styles = getSampleStyleSheet()

        # Title
        title_style = ParagraphStyle('TitleStyle', parent=styles['Title'],
                                     alignment=1, fontSize=16, spaceAfter=8,
                                     fontName='Helvetica-Bold')
        elements.append(Paragraph(title, title_style))
        elements.append(Spacer(1, 4))

        # Details – with formatted date
        detail_data = [
            [emp_label, schedule_data.get("employee_name", "")],
            [pos_label, schedule_data.get("position", "")],
            [date_label, format_date_pdf(schedule_data.get("date", ""))]
        ]
        detail_table = Table(detail_data, colWidths=[2.0*inch, 5.5*inch])
        detail_table.setStyle(TableStyle([
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,0), (-1,-1), 10),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LEFTPADDING', (0,0), (-1,-1), 2),
            ('RIGHTPADDING', (0,0), (-1,-1), 2),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
        ]))
        elements.append(detail_table)
        elements.append(Spacer(1, 10))

        # Schedule table – FILL WITH DATA
        schedule = schedule_data.get("schedule", [])
        table_data = [[time_header, act_header, desc_header]]
        for entry in schedule:
            table_data.append([
                entry.get("slot", ""),
                entry.get("activity", ""),
                entry.get("description", "")
            ])

        col_widths = [1.2*inch, 2.2*inch, 3.9*inch]
        schedule_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        schedule_table.setStyle(TableStyle([
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,0), (-1,-1), 8),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('GRID', (0,0), (-1,-1), 0.5, colors.black),
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#CCCCCC")),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 9),
            ('ALIGN', (0,0), (-1,0), 'CENTER'),
            ('WORDWRAP', (0,0), (-1,-1), True),
        ]))
        elements.append(schedule_table)

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    except Exception as e:
        print(f"⚠️ Template fallback error: {e} – using basic fallback")
        return _create_basic_fallback_pdf(schedule_data)

# -------- Basic fallback (last resort) --------
def _create_basic_fallback_pdf(schedule_data):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            rightMargin=15, leftMargin=15,
                            topMargin=15, bottomMargin=15)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('TitleStyle', parent=styles['Title'],
                                 alignment=1, fontSize=16, spaceAfter=8,
                                 fontName='Helvetica-Bold')
    elements.append(Paragraph("EOD REPORT", title_style))
    elements.append(Spacer(1, 4))

    # Details – with formatted date
    detail_data = [
        ["Name Of Employee", schedule_data.get("employee_name", "")],
        ["Position", schedule_data.get("position", "")],
        ["Date", format_date_pdf(schedule_data.get("date", ""))]
    ]
    detail_table = Table(detail_data, colWidths=[2.0*inch, 5.5*inch])
    detail_table.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('LEFTPADDING', (0,0), (-1,-1), 2),
        ('RIGHTPADDING', (0,0), (-1,-1), 2),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
    ]))
    elements.append(detail_table)
    elements.append(Spacer(1, 10))

    schedule = schedule_data.get("schedule", [])
    table_data = [["Time", "Activity", "Description"]]
    for entry in schedule:
        table_data.append([
            entry.get("slot", ""),
            entry.get("activity", ""),
            entry.get("description", "")
        ])

    col_widths = [1.2*inch, 2.2*inch, 3.9*inch]
    schedule_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    schedule_table.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('GRID', (0,0), (-1,-1), 0.5, colors.black),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#CCCCCC")),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 9),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('WORDWRAP', (0,0), (-1,-1), True),
    ]))
    elements.append(schedule_table)
    doc.build(elements)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes